"""Export a Red Tide map workbook to the JSON files the game loads.
Usage: python export_json.py RedTide_SunkenShip_data.xlsx redtide_maps/ship
Writes species.json, attacks.json, diet.json, flora.json, spawn.json, tide_curve.json, tunables.json,
blockout.json (zones, links, points of interest and any map-specific tables), boss.json, faction.json.
Run the workbook through recalc (LibreOffice) or open+save it in Excel first so formula cells have values."""
import sys, json, os
from openpyxl import load_workbook
src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
wb = load_workbook(src, data_only=True)
def rows(ws, header_row=1, stop_blank=True):
    hdr = [c.value for c in ws[header_row]]
    n = max(i for i, h in enumerate(hdr) if h is not None) + 1
    hdr = hdr[:n]; res = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        r = r[:n]
        if r[0] is None:
            if stop_blank: break
            continue
        res.append(dict(zip(hdr, r)))
    return res
def dump(name, obj):
    with open(os.path.join(out, name), "w") as f: json.dump(obj, f, indent=1)
dump("species.json", rows(wb["Species"]))
dump("attacks.json", rows(wb["Attacks"]))
ws = wb["Diet"]; hdr = [c.value for c in ws[1]]; foods = hdr[1:-1]; diet = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None: break
    diet[r[0]] = {f: v for f, v in zip(foods, r[1:-1]) if v}
dump("diet.json", diet)
dump("flora.json", rows(wb["Flora"]))
dump("spawn.json", rows(wb["Spawn"]))
ws = wb["TideCurve"]
tun = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(2, 18) if ws.cell(row=i, column=1).value}
dump("tunables.json", tun); dump("tide_curve.json", rows(ws, header_row=18))
# Blockout: a sheet of several titled tables; split on title rows (single non-empty cell, followed by a header row)
ws = wb["Blockout"]; tables = {}; grid = [[c for c in r] for r in ws.iter_rows(values_only=True)]
i = 0
while i < len(grid):
    row = grid[i]; filled = [c for c in row if c is not None]
    if len(filled) == 1 and i + 1 < len(grid) and sum(c is not None for c in grid[i + 1]) > 1:
        title = str(filled[0]).split(" (")[0].strip().lower().replace(" ", "_").replace("'", "")
        hdr = [c for c in grid[i + 1] if c is not None]; items = []; j = i + 2
        while j < len(grid) and grid[j][0] is not None:
            items.append(dict(zip(hdr, grid[j][:len(hdr)]))); j += 1
        tables[title] = items; i = j
    elif len(filled) == 1 and str(filled[0]).startswith("Design notes"):
        notes = []; j = i + 1
        while j < len(grid) and grid[j][0] is not None: notes.append(grid[j][0]); j += 1
        tables["design_notes"] = notes; i = j
    else: i += 1
dump("blockout.json", tables)
dump("boss.json", [ws.cell(row=r, column=1).value for r in range(1, wb["Boss"].max_row + 1) if wb["Boss"].cell(row=r, column=1).value] if (ws := wb["Boss"]) else [])
if "Faction" in wb.sheetnames:
    ws = wb["Faction"]; fa = {"name": ws["A1"].value, "intro": ws["A2"].value, "units": rows(ws, header_row=4), "tactics": [], "barks": []}
    mode = None
    for r in ws.iter_rows(min_row=5, values_only=True):
        v = r[0]
        if v == "Squad tactics": mode = "tactics"; continue
        if isinstance(v, str) and v.startswith("Barks"): mode = "barks"; continue
        if mode and v and not (isinstance(v, str) and v.startswith("Enemy HP")): fa[mode].append(v)
    dump("faction.json", fa)
print("exported", out, sorted(os.listdir(out)))
