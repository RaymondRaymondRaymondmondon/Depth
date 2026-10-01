import sys, importlib
mod=importlib.import_module(sys.argv[1])
from species_ship import SIZE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
wb=Workbook(); F=Font(name="Arial",size=10); H=Font(name="Arial",size=10,bold=True); B=Font(name="Arial",size=10,color="0000FF"); FILL=PatternFill("solid",fgColor="DDEBF7")
def hdr(ws,row,vals):
    for i,v in enumerate(vals,1):
        c=ws.cell(row=row,column=i,value=v); c.font=H; c.fill=FILL
def fit(ws,w=14):
    for i in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(i)].width=w
S,D,FOODS,FLORA,ZONES,LINKS,POI,SPAWN,BOSS=mod.S,mod.D,mod.FOODS,mod.FLORA,mod.ZONES,mod.LINKS,mod.POI,mod.SPAWN,mod.BOSS
ws=wb.active; ws.title="README"
rows=[f"Red Tide — {mod.MAP}: map data workbook","",
"Sheets: Species (blue = tunable inputs), Attacks, Diet (predator rows x food columns, weights sum to 1), Flora, Spawn, TideCurve (formulas from the Tunables block), Blockout (zones, links, points of interest in metres), Boss, Faction.",
"Claude Code loads these as redtide_maps/"+mod.SLUG+"/*.json. Change blue cells only; formulas recompute.",
"Rules: HP = base x 1.08^(tide-1); bounty = base x 1.10^(tide-1); blood on death = size x 20; blood/s wounded = size x 2 x wound fraction.",
"Sizes 1-6 map to HP 20/60/250/900/3000/15000 and bounty 40/80/180/400/900/1500 (design doc, Shared beast rules).",
"All values are design targets from the Red Tide design document (2026-09-29); tune with --eco-sim and --redtide-sim."]+getattr(mod,"README_EXTRA",[])
for i,r in enumerate(rows,1): ws.cell(row=i,column=1,value=r).font=H if i==1 else F
ws.column_dimensions["A"].width=150
ws=wb.create_sheet("Species")
cols=["id","name","class","size","tier","tags","archetype","social","group_size","home_zone","home_flora","defend_radius_m","blood_threshold","aggression","fear","curiosity","sight_m","scent_m","hearing_m","electro_m","weak_point","armor_front","hp_base","bounty_base","blood_death","blood_per_s_wounded","speed_mps","turn_deg_s","drop_pct","hp_tide10","bounty_tide10","special"]
hdr(ws,1,cols)
for r,s in enumerate(S,2):
    (name,cls,size,tier,tags,arch,social,grp,hz,hf,dr,thr,ag,fe,cu,si,sc,he,el,wp,af,att,special,drop)=s
    sz=SIZE[size]
    vals=[r-1,name,cls,size,tier,tags,arch,social,grp,hz,hf,dr if dr!="" else 0,thr,ag,fe,cu,si,sc,he,el,wp,af,sz["hp"],sz["bounty"],None,None,sz["speed"],sz["turn"],drop,None,None,special]
    for c,v in enumerate(vals,1): ws.cell(row=r,column=c,value=v).font=F
    for c in (4,5,12,13,14,15,16,17,18,19,20,22,23,24,27,28,29): ws.cell(row=r,column=c).font=B
    ws.cell(row=r,column=25,value=f"=D{r}*20"); ws.cell(row=r,column=26,value=f"=D{r}*2")
    ws.cell(row=r,column=30,value=f"=ROUND(W{r}*1.08^9,0)"); ws.cell(row=r,column=31,value=f"=ROUND(X{r}*1.1^9,0)")
fit(ws,13); ws.column_dimensions["B"].width=24; ws.column_dimensions["AF"].width=70; ws.freeze_panes="C2"
ws.cell(row=len(S)+3,column=1,value="Blue = tunable input. Columns Y, Z, AD, AE are formulas.").font=F
ws=wb.create_sheet("Attacks"); hdr(ws,1,["beast","attack","damage","windup_s","cooldown_s","range_m","tell (what the player sees)","effect"])
r=2
for s in S:
    for a in s[21]:
        for c,v in enumerate([s[0]]+list(a),1): ws.cell(row=r,column=c,value=v).font=B if c in (3,4,5,6) else F
        r+=1
ws.cell(row=r+1,column=1,value="Beasts with no rows never attack. Damage is vs 100 HP divers at tide 1, x1.05 per tide (TideCurve), capped x2.5.").font=F
fit(ws,14); ws.column_dimensions["G"].width=36; ws.column_dimensions["H"].width=60
ws=wb.create_sheet("Diet")
names=[s[0] for s in S]; foods=names+FOODS
hdr(ws,1,["predator \\ food"]+foods+["row_sum"])
for r,p in enumerate(names,2):
    ws.cell(row=r,column=1,value=p).font=H
    for c,f in enumerate(foods,2):
        w=D.get(p,{}).get(f,0)
        if w: ws.cell(row=r,column=c,value=w).font=B
    ws.cell(row=r,column=len(foods)+2,value=f"=SUM(B{r}:{get_column_letter(len(foods)+1)}{r})").font=F
ws.cell(row=len(names)+3,column=1,value="Preference shares; each row sums to 1 (0 for non-feeders). Blank = never eats. Predators also obey the size rule (prey <= 60% of size unless pack/school).").font=F
ws.freeze_panes="B2"; ws.column_dimensions["A"].width=24
for i in range(2,len(foods)+3): ws.column_dimensions[get_column_letter(i)].width=6
missing=set(f for d in D.values() for f in d)-set(foods); assert not missing, missing
ws=wb.create_sheet("Flora"); hdr(ws,1,["flora","type","zones","effect","contact_damage","grazed_by","regrowth_rate_per_s","regrow_delay_s"])
for r,f in enumerate(FLORA,2):
    for c,v in enumerate(f,1): ws.cell(row=r,column=c,value=v).font=B if c in (5,7,8) else F
fit(ws,16); ws.column_dimensions["D"].width=64; ws.column_dimensions["F"].width=44
ws=wb.create_sheet("Spawn"); hdr(ws,1,["zone","species","count_tide1","respawn_s","capacity_mult_per_tide","count_tide10"])
for r,(z,sp,cnt,rs) in enumerate(SPAWN,2):
    assert sp in names, sp
    for c,v in enumerate([z,sp,cnt,rs,1.03],1): ws.cell(row=r,column=c,value=v).font=B if c in (3,4,5) else F
    ws.cell(row=r,column=6,value=f"=ROUND(C{r}*E{r}^9,0)").font=F
ws.cell(row=len(SPAWN)+3,column=1,value="respawn_s: seconds to replace one killed individual while below capacity; 0 = placed once (boss, sessile, ambush nests). Swarms and colonies count as 1 agent.").font=F
fit(ws,22)
ws=wb.create_sheet("TideCurve")
tun=[("base_quota_4p",12),("quota_growth",1.15),("hp_growth",1.08),("bounty_growth",1.10),("dmg_growth",1.05),("alarm_base",100),("alarm_tide_factor",0.1),("spawn_chance_base",0.10),("spawn_chance_per_tide",0.03),("spawn_chance_cap",0.70),("blood_decay_early",0.02),("blood_decay_mid",0.01),("blood_decay_late",0.005),("calm_seconds",20)]
ws.cell(row=1,column=1,value="Tunables").font=H
for i,(k,v) in enumerate(tun,2): ws.cell(row=i,column=1,value=k).font=F; ws.cell(row=i,column=2,value=v).font=B
hdr(ws,18,["tide","kill_quota_4p","hp_mult","bounty_mult","beast_dmg_mult","alarm_threshold","enemy_spawn_chance","blood_decay_per_s","hunt","apex_wander","boss_may_appear","tide_bonus_scrip"])
for t in range(1,41):
    r=18+t; ws.cell(row=r,column=1,value=t)
    for c,f in enumerate([f"=ROUND($B$2*$B$3^(A{r}-1),0)",f"=ROUND($B$4^(A{r}-1),3)",f"=ROUND($B$5^(A{r}-1),3)",f"=ROUND(MIN(2.5,$B$6^(A{r}-1)),3)",f"=ROUND($B$7/(1+$B$8*A{r}),1)",f"=IF(A{r}<4,0,MIN($B$11,$B$9+$B$10*A{r}))",f"=IF(A{r}<10,$B$12,IF(A{r}<20,$B$13,$B$14))",f'=IF(A{r}<5,"none",IF(MOD(A{r},5)=0,IF(A{r}>=10,"Faction or Predator","Faction"),"none"))',f'=IF(A{r}>=10,"yes","no")',f'=IF(A{r}>=12,"yes",IF(A{r}>=10,"Predator Hunt only","no"))',f"=100*A{r}"],2):
        ws.cell(row=r,column=c,value=f)
    for c in range(1,13): ws.cell(row=r,column=c).font=F
ws.cell(row=60,column=1,value="Quota scales 0.4/0.6/0.8/1.0 for 1-4 players.").font=F
fit(ws,18); ws.column_dimensions["A"].width=24
ws=wb.create_sheet("Blockout")
ws.cell(row=1,column=1,value="Zones (metres)").font=H
hdr(ws,2,getattr(mod,"ZONE_HEADER",["zone","level","x","y","w","h","door_cost_scrip","depth_m","notes"]))
for r,z in enumerate(ZONES,3):
    for c,v in enumerate(z,1): ws.cell(row=r,column=c,value=v).font=B if c in (3,4,5,6,7) else F
r0=len(ZONES)+5; ws.cell(row=r0-1,column=1,value="Links").font=H; hdr(ws,r0,["from","to","cost_scrip","passage"])
for r,l in enumerate(LINKS,r0+1):
    for c,v in enumerate(l,1): ws.cell(row=r,column=c,value=v).font=F
r1=r0+len(LINKS)+3
if hasattr(mod,"STREAMS"):
    ws.cell(row=r1-1,column=1,value="Slipstreams (one-way, 8 m/s, no firing inside)").font=H; hdr(ws,r1,["id","from","to","length_m","ride_s","hazard"])
    for r,l in enumerate(mod.STREAMS,r1+1):
        for c,v in enumerate(l,1): ws.cell(row=r,column=c,value=v).font=F
    r1=r1+len(mod.STREAMS)+3
for attr,title,cols in [("BUILDINGS","Building types (the city generator's catalogue)",["type","district","count","footprint_m","floors","hosts (habitat rule)","contents / loot","placement rule"]),("AQUEDUCT","Aqueduct (blood-flow network)",["segment","from","to","length_m","blood_transit_s","note"]),("GRATES","Cistern grates (the Wyrm's and the congers' exits)",["id","district","x","y","note"])]:
    if hasattr(mod,attr):
        ws.cell(row=r1-1,column=1,value=title).font=H; hdr(ws,r1,cols)
        for r,l in enumerate(getattr(mod,attr),r1+1):
            for c,v in enumerate(l,1): ws.cell(row=r,column=c,value=v).font=F
        r1=r1+len(getattr(mod,attr))+3
if hasattr(mod,"CORRIDORS"):
    ws.cell(row=r1-1,column=1,value="Corridors").font=H; hdr(ws,r1,["id","zone","length_m","width_m","eel_holes","reacher_gate","note"])
    for r,l in enumerate(mod.CORRIDORS,r1+1):
        for c,v in enumerate(l,1): ws.cell(row=r,column=c,value=v).font=F
    r1=r1+len(mod.CORRIDORS)+3
ws.cell(row=r1-1,column=1,value="Points of interest").font=H; hdr(ws,r1,["name","type","zone","x","y"])
for r,p in enumerate(POI,r1+1):
    for c,v in enumerate(p,1): ws.cell(row=r,column=c,value=v).font=F
r2=r1+len(POI)+3; ws.cell(row=r2-1,column=1,value="Design notes").font=H
for i,nn in enumerate(getattr(mod,"NOTES",[])): ws.cell(row=r2+i,column=1,value=nn).font=F
fit(ws,16); ws.column_dimensions["I"].width=80; ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=30; ws.column_dimensions["C"].width=30
ws=wb.create_sheet("Boss")
for r,b in enumerate(BOSS,1): ws.cell(row=r,column=1,value=b).font=F
ws.column_dimensions["A"].width=170
if hasattr(mod,"FACTION"):
    fa=mod.FACTION; ws=wb.create_sheet("Faction")
    ws.cell(row=1,column=1,value=fa["name"]).font=H; ws.cell(row=2,column=1,value=fa["intro"]).font=F
    hdr(ws,4,["unit","hp_tide1","weapon","damage","speed","role","tell","loot"])
    for r,u in enumerate(fa["units"],5):
        for c,v in enumerate(u,1): ws.cell(row=r,column=c,value=v).font=B if c==2 else F
    r=5+len(fa["units"])+1; ws.cell(row=r,column=1,value="Squad tactics").font=H
    for i,t in enumerate(fa["tactics"],1): ws.cell(row=r+i,column=1,value=f"{i}. {t}").font=F
    r=r+len(fa["tactics"])+2; ws.cell(row=r,column=1,value="Barks (formant synth, subtitled)").font=H
    for i,t in enumerate(fa["barks"],1): ws.cell(row=r+i,column=1,value=t).font=F
    ws.cell(row=r+len(fa["barks"])+2,column=1,value="Enemy HP x1.08 per tide; weapons upgrade every 5 tides (Forged at 15+).").font=F
    fit(ws,18); ws.column_dimensions["A"].width=40; ws.column_dimensions["G"].width=50; ws.column_dimensions["F"].width=44
out="/mnt/user-data/outputs/"+mod.FILE; wb.save(out); print("saved",out,len(S),"species")
